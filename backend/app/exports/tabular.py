"""CSV and XLSX exports."""

from __future__ import annotations

import csv
import io
from typing import Any

COLUMNS = [
    ("wbs_code", "WBS"),
    ("id", "Task ID"),
    ("label", "Task Name"),
    ("work_package", "Work Package"),
    ("storey_name", "Storey"),
    ("zone_name", "Zone"),
    ("ifc_class", "IFC Class"),
    ("predefined_type", "Predefined Type"),
    ("type_name", "Type"),
    ("material", "Material"),
    ("element_count", "Elements"),
    ("quantity", "Quantity"),
    ("unit", "Unit"),
    ("quantity_source", "Quantity Source"),
    ("duration_days", "Duration (d)"),
    ("crew", "Crew"),
    ("rate_id", "Rate"),
    ("confidence", "Confidence"),
    ("start_date", "Start"),
    ("finish_date", "Finish"),
    ("total_float", "Total Float"),
    ("free_float", "Free Float"),
    ("is_critical", "Critical"),
]


def _predecessor_text(task: dict[str, Any]) -> str:
    parts = []
    for link in task.get("predecessors") or []:
        lag = int(link.get("lag") or 0)
        suffix = f"{lag:+d}" if lag else ""
        parts.append(f"{link.get('id')}{link.get('type', 'FS')}{suffix}")
    return ", ".join(parts)


# Appended when the schedule is being tracked against progress.
PROGRESS_COLUMNS = [
    ("status", "Status"),
    ("percent_complete", "% Complete"),
    ("quantity_placed", "Qty Placed"),
    ("actual_start", "Actual Start"),
    ("actual_finish", "Actual Finish"),
    ("baseline_start", "Baseline Start"),
    ("baseline_finish", "Baseline Finish"),
    ("forecast_finish", "Forecast Finish"),
    ("finish_variance_days", "Finish Variance (d)"),
    ("schedule_flag", "Flag"),
    ("delay_cause", "Delay Cause"),
]


def _progress_for(schedule: dict[str, Any]) -> dict[str, dict[str, Any]] | None:
    progress = schedule.get("progress")
    return progress.get("tasks") if isinstance(progress, dict) else None


def _row(task: dict[str, Any], progress: dict[str, dict[str, Any]] | None = None) -> list[Any]:
    row = [task.get(key) for key, _ in COLUMNS]
    if progress is not None:
        state = progress.get(task.get("id"), {})
        row.extend(state.get(key) for key, _ in PROGRESS_COLUMNS)
    row.append(_predecessor_text(task))
    row.append(";".join(task.get("element_ids") or []))
    return row


def _header(tracked: bool = False) -> list[str]:
    extra = [title for _, title in PROGRESS_COLUMNS] if tracked else []
    return [title for _, title in COLUMNS] + extra + ["Predecessors", "Source GlobalIds"]


def to_csv(schedule: dict[str, Any], project_name: str = "Schedule") -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    progress = _progress_for(schedule)
    writer.writerow(_header(progress is not None))
    for task in schedule.get("tasks") or []:
        writer.writerow(_row(task, progress))
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(schedule: dict[str, Any], project_name: str = "Schedule") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Schedule"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")
    critical_fill = PatternFill("solid", fgColor="FEE2E2")

    progress = _progress_for(schedule)
    behind_fill = PatternFill("solid", fgColor="FEF3C7")
    header = _header(progress is not None)
    sheet.append(header)
    for column in range(1, len(header) + 1):
        cell = sheet.cell(row=1, column=column)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for task in schedule.get("tasks") or []:
        sheet.append(_row(task, progress))
        behind = progress is not None and (
            progress.get(task.get("id"), {}).get("schedule_flag") == "behind"
        )
        # Late beats critical: it is the more urgent of the two to act on.
        fill = behind_fill if behind else critical_fill if task.get("is_critical") else None
        if fill is not None:
            for column in range(1, len(header) + 1):
                sheet.cell(row=sheet.max_row, column=column).fill = fill

    widths = [10, 18, 46, 16, 14, 12, 18, 18, 22, 22, 10, 12, 8, 16, 12, 8, 18, 12,
              12, 12, 12, 12, 10]
    if progress is not None:
        widths += [12, 11, 11, 13, 13, 14, 14, 15, 12, 11, 18]
    widths += [30, 40]
    for index, width in enumerate(widths[: len(header)], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"

    report_sheet = workbook.create_sheet("Run report")
    report_sheet.append(["Metric", "Value"])
    report_sheet.cell(row=1, column=1).font = Font(bold=True)
    report_sheet.cell(row=1, column=2).font = Font(bold=True)
    for key, value in _flatten(schedule.get("report") or {}):
        report_sheet.append([key, value])
    report_sheet.column_dimensions["A"].width = 46
    report_sheet.column_dimensions["B"].width = 60

    summary = (schedule.get("progress") or {}).get("summary")
    if summary:
        progress_sheet = workbook.create_sheet("Progress", 1)
        progress_sheet.append(["Metric", "Value"])
        for column in (1, 2):
            progress_sheet.cell(row=1, column=column).font = Font(bold=True)
        for key, value in _flatten(summary):
            progress_sheet.append([key, value])
        progress_sheet.column_dimensions["A"].width = 46
        progress_sheet.column_dimensions["B"].width = 60

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _flatten(data: Any, prefix: str = "") -> list[tuple[str, Any]]:
    rows: list[tuple[str, Any]] = []
    if isinstance(data, dict):
        for key, value in data.items():
            rows.extend(_flatten(value, f"{prefix}.{key}" if prefix else str(key)))
    elif isinstance(data, list):
        rows.append((prefix, ", ".join(str(item) for item in data[:20])))
    else:
        rows.append((prefix, data))
    return rows
