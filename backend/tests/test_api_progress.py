"""API: baselines, the data date and progress reporting."""

from __future__ import annotations

import pytest


@pytest.fixture
def ordered_tasks(api_client, scheduled):
    tasks = api_client.get(f"/api/projects/{scheduled}/schedule").json()["tasks"]
    return sorted(tasks, key=lambda t: (t["early_start_offset"], t["id"]))


@pytest.fixture
def baselined(api_client, scheduled):
    response = api_client.post(f"/api/projects/{scheduled}/baselines", json={"name": "Contract"})
    assert response.status_code == 200
    api_client.put(f"/api/projects/{scheduled}/data-date", json={"data_date": "2026-09-15"})
    return scheduled


def report(api_client, project, entries, reported_on=None):
    body = {"entries": entries}
    if reported_on:
        body["reported_on"] = reported_on
    return api_client.post(f"/api/projects/{project}/progress", json=body)


# --- baselines --------------------------------------------------------------


def test_baseline_snapshot_and_listing(api_client, scheduled):
    created = api_client.post(f"/api/projects/{scheduled}/baselines", json={}).json()
    assert created["name"] == "Baseline 1"
    assert created["is_current"] is True
    assert created["task_count"] > 0
    assert created["finish_date"]

    second = api_client.post(
        f"/api/projects/{scheduled}/baselines", json={"name": "Re-baseline"}
    ).json()
    listed = api_client.get(f"/api/projects/{scheduled}/baselines").json()["baselines"]
    assert [b["name"] for b in listed] == ["Re-baseline", "Baseline 1"]
    assert [b["is_current"] for b in listed] == [True, False]

    api_client.post(f"/api/projects/{scheduled}/baselines/{created['id']}/activate")
    listed = api_client.get(f"/api/projects/{scheduled}/baselines").json()["baselines"]
    assert {b["id"]: b["is_current"] for b in listed} == {
        second["id"]: False,
        created["id"]: True,
    }


def test_deleting_the_current_baseline_promotes_the_newest(api_client, scheduled):
    first = api_client.post(f"/api/projects/{scheduled}/baselines", json={}).json()
    second = api_client.post(f"/api/projects/{scheduled}/baselines", json={}).json()
    api_client.delete(f"/api/projects/{scheduled}/baselines/{second['id']}")
    listed = api_client.get(f"/api/projects/{scheduled}/baselines").json()["baselines"]
    assert [(b["id"], b["is_current"]) for b in listed] == [(first["id"], True)]


def test_baseline_survives_edits_to_the_plan(api_client, baselined, ordered_tasks):
    """The whole point of a baseline: editing the plan must not move it."""
    first = ordered_tasks[0]
    api_client.patch(
        f"/api/projects/{baselined}/tasks/{first['id']}",
        json={"duration_days": first["duration_days"] + 5},
    )
    progress = api_client.get(f"/api/projects/{baselined}/progress").json()
    row = next(t for t in progress["tasks"] if t["id"] == first["id"])
    assert row["baseline_duration"] == first["duration_days"]
    assert row["baseline_finish"] == first["finish_date"]
    assert progress["summary"]["finish_variance_days"] > 0


def test_baseline_needs_a_schedule(api_client, project):
    assert api_client.post(f"/api/projects/{project}/baselines", json={}).status_code == 404


# --- data date ----------------------------------------------------------------


def test_data_date_is_stored_and_validated(api_client, scheduled):
    response = api_client.put(
        f"/api/projects/{scheduled}/data-date", json={"data_date": "2026-09-10"}
    )
    assert response.json() == {"data_date": "2026-09-10"}
    assert api_client.get(f"/api/projects/{scheduled}").json()["data_date"] == "2026-09-10"

    bad = api_client.put(f"/api/projects/{scheduled}/data-date", json={"data_date": "soon"})
    assert bad.status_code == 422


def test_data_date_defaults_to_today(api_client, scheduled):
    payload = api_client.get(f"/api/projects/{scheduled}/progress").json()
    assert payload["data_date_is_default"] is True


# --- reporting ------------------------------------------------------------------


def test_report_and_read_back(api_client, baselined, ordered_tasks):
    first = ordered_tasks[0]
    response = report(
        api_client,
        baselined,
        [{"task_id": first["id"], "percent_complete": 100, "actual_start": "2026-09-01",
          "actual_finish": "2026-09-07"}],
    )
    assert response.status_code == 200
    payload = response.json()
    row = next(t for t in payload["tasks"] if t["id"] == first["id"])
    assert row["status"] == "complete"
    assert row["actual_finish"] == "2026-09-07"
    assert row["schedule_flag"] == "complete"
    assert payload["summary"]["variance_basis"] == "baseline"
    assert payload["baseline"]["name"] == "Contract"

    # Nothing was written back onto the plan.
    plan = api_client.get(f"/api/projects/{baselined}/schedule").json()
    plan_row = next(t for t in plan["tasks"] if t["id"] == first["id"])
    assert "actual_finish" not in plan_row


def test_partial_reports_merge_instead_of_overwriting(api_client, baselined, ordered_tasks):
    target = next(t for t in ordered_tasks if t["quantity"])
    report(
        api_client, baselined,
        [{"task_id": target["id"], "actual_start": "2026-09-02"}],
        reported_on="2026-09-02",
    )
    payload = report(
        api_client, baselined,
        [{"task_id": target["id"], "quantity_placed": target["quantity"] / 4}],
        reported_on="2026-09-09",
    ).json()
    row = next(t for t in payload["tasks"] if t["id"] == target["id"])
    # The later report said nothing about the start date; it must survive.
    assert row["actual_start"] == "2026-09-02"
    assert row["percent_complete"] == 25.0
    assert row["status"] == "in_progress"


def test_later_reports_override_earlier_values(api_client, baselined, ordered_tasks):
    target = ordered_tasks[0]
    report(api_client, baselined, [{"task_id": target["id"], "percent_complete": 20}],
           reported_on="2026-09-03")
    payload = report(api_client, baselined, [{"task_id": target["id"], "percent_complete": 60}],
                     reported_on="2026-09-10").json()
    row = next(t for t in payload["tasks"] if t["id"] == target["id"])
    assert row["percent_complete"] == 60.0


def test_history_is_kept(api_client, baselined, ordered_tasks):
    target = ordered_tasks[0]
    for day, pct in (("2026-09-02", 10), ("2026-09-04", 40), ("2026-09-08", 70)):
        report(api_client, baselined, [{"task_id": target["id"], "percent_complete": pct}],
               reported_on=day)
    history = api_client.get(
        f"/api/projects/{baselined}/progress/history", params={"task_id": target["id"]}
    ).json()["updates"]
    assert [u["percent_complete"] for u in history] == [70, 40, 10]


def test_clearing_a_task_returns_it_to_not_started(api_client, baselined, ordered_tasks):
    target = ordered_tasks[0]
    report(api_client, baselined, [{"task_id": target["id"], "percent_complete": 50}])
    payload = api_client.delete(f"/api/projects/{baselined}/progress/{target['id']}").json()
    assert payload["removed_updates"] == 1
    row = next(t for t in payload["tasks"] if t["id"] == target["id"])
    assert row["status"] == "not_started"


def test_slip_is_visible_in_the_summary(api_client, baselined, ordered_tasks):
    first = ordered_tasks[0]
    before = api_client.get(f"/api/projects/{baselined}/progress").json()["summary"]
    payload = report(
        api_client, baselined,
        [{"task_id": first["id"], "actual_start": "2026-09-01", "actual_finish": "2026-09-14"}],
    ).json()
    after = payload["summary"]
    assert after["forecast_finish"] > before["reference_finish"]
    assert after["finish_variance_days"] > 0
    assert after["critical_behind"]


# --- validation ------------------------------------------------------------------


def test_unknown_task_ids_are_rejected(api_client, baselined):
    response = report(api_client, baselined, [{"task_id": "nope", "percent_complete": 10}])
    assert response.status_code == 422
    assert "nope" in response.json()["detail"]


def test_actuals_after_the_data_date_are_rejected(api_client, baselined, ordered_tasks):
    response = report(
        api_client, baselined,
        [{"task_id": ordered_tasks[0]["id"], "actual_start": "2026-12-01"}],
    )
    assert response.status_code == 422
    assert "data date" in response.json()["detail"]


def test_finish_before_start_is_rejected(api_client, baselined, ordered_tasks):
    response = report(
        api_client, baselined,
        [{"task_id": ordered_tasks[0]["id"], "actual_start": "2026-09-08",
          "actual_finish": "2026-09-02"}],
    )
    assert response.status_code == 422


def test_empty_updates_are_rejected(api_client, baselined, ordered_tasks):
    response = report(api_client, baselined, [{"task_id": ordered_tasks[0]["id"]}])
    assert response.status_code == 422
    assert "reports nothing" in response.json()["detail"]


def test_out_of_range_percent_is_rejected(api_client, baselined, ordered_tasks):
    response = report(
        api_client, baselined, [{"task_id": ordered_tasks[0]["id"], "percent_complete": 140}]
    )
    assert response.status_code == 422


def test_bad_dates_are_rejected(api_client, baselined, ordered_tasks):
    response = report(
        api_client, baselined, [{"task_id": ordered_tasks[0]["id"], "actual_start": "last week"}]
    )
    assert response.status_code == 422


def test_progress_needs_a_schedule(api_client, project):
    assert api_client.get(f"/api/projects/{project}/progress").status_code == 404


def test_regenerating_orphans_old_progress_with_a_warning(api_client, baselined, ordered_tasks):
    report(api_client, baselined, [{"task_id": ordered_tasks[0]["id"], "percent_complete": 30}])

    from tests.conftest import wait_for_job

    job = api_client.post(
        f"/api/projects/{baselined}/schedule", json={"level": "L1", "start_date": "2026-09-01"}
    ).json()["job_id"]
    assert wait_for_job(api_client, job)["status"] == "done"
    payload = api_client.get(f"/api/projects/{baselined}/progress").json()
    assert payload["warnings"]


def test_deleting_a_project_removes_its_progress_and_baselines(api_client, baselined,
                                                              ordered_tasks):
    report(api_client, baselined, [{"task_id": ordered_tasks[0]["id"], "percent_complete": 30}])
    assert api_client.delete(f"/api/projects/{baselined}").status_code == 200

    from app.db import BaselineRecord, ProgressUpdate, session_scope

    with session_scope() as session:
        assert session.query(ProgressUpdate).filter_by(project_id=baselined).count() == 0
        assert session.query(BaselineRecord).filter_by(project_id=baselined).count() == 0


def test_old_databases_gain_new_columns(tmp_path, monkeypatch):
    """A database created before these columns existed must still work."""
    import sqlite3

    db_path = tmp_path / "legacy.db"
    connection = sqlite3.connect(db_path)
    connection.execute(
        "CREATE TABLE projects (id INTEGER PRIMARY KEY, name VARCHAR(200) NOT NULL, "
        "status VARCHAR(40))"
    )
    connection.execute("INSERT INTO projects (id, name, status) VALUES (1, 'Old', 'parsed')")
    connection.commit()
    connection.close()

    monkeypatch.setenv("IFCSCHED_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("IFCSCHED_DATABASE_URL", f"sqlite:///{db_path}")
    from app import config as config_module
    from app import db as db_module

    config_module.get_settings.cache_clear()
    db_module.reset_engine()
    try:
        db_module.init_db()
        columns = {
            row[1] for row in sqlite3.connect(db_path).execute("PRAGMA table_info(projects)")
        }
        assert {"data_date", "geometry_status", "profile"} <= columns
        with db_module.session_scope() as session:
            assert session.get(db_module.Project, 1).name == "Old"
    finally:
        db_module.reset_engine()
        config_module.get_settings.cache_clear()


# --- exports ------------------------------------------------------------------


@pytest.fixture
def tracked(api_client, baselined, ordered_tasks):
    """A schedule with one finished task and one half-done task.

    The half-done one must last several days: in-progress work always keeps at
    least a day remaining, so on a one-day task "remaining" equals "planned".
    """
    first = ordered_tasks[0]
    second = next(
        t for t in ordered_tasks if t["id"] != first["id"] and t["duration_days"] >= 4
    )
    report(
        api_client,
        baselined,
        [
            {"task_id": first["id"], "actual_start": "2026-09-01", "actual_finish": "2026-09-07"},
            {"task_id": second["id"], "percent_complete": 50, "actual_start": "2026-09-02"},
        ],
    )
    return baselined, first, second


def test_untracked_exports_carry_no_progress(api_client, scheduled):
    import csv
    import io
    import json

    text = api_client.get(f"/api/projects/{scheduled}/export/csv").content.decode("utf-8-sig")
    assert "% Complete" not in next(csv.reader(io.StringIO(text)))
    package = json.loads(api_client.get(f"/api/projects/{scheduled}/export/json").content)
    assert package["progress"] is None
    assert all(task["progress"] is None for task in package["tasks"])
    xer = api_client.get(f"/api/projects/{scheduled}/export/xer").text
    assert "phys_complete_pct" not in xer


def test_csv_and_xlsx_carry_progress(api_client, tracked):
    import csv
    import io

    from openpyxl import load_workbook

    project, first, second = tracked
    text = api_client.get(f"/api/projects/{project}/export/csv").content.decode("utf-8-sig")
    rows = {row["Task ID"]: row for row in csv.DictReader(io.StringIO(text))}
    assert rows[first["id"]]["Status"] == "complete"
    assert rows[first["id"]]["Actual Finish"] == "2026-09-07"
    assert rows[second["id"]]["% Complete"] == "50.0"
    assert rows[first["id"]]["Baseline Finish"] == first["finish_date"]
    # GlobalIds are still the last column, whatever else was added.
    assert list(rows[first["id"]])[-1] == "Source GlobalIds"

    xlsx = api_client.get(f"/api/projects/{project}/export/xlsx").content
    workbook = load_workbook(io.BytesIO(xlsx))
    assert workbook.sheetnames == ["Schedule", "Progress", "Run report"]
    metrics = {row[0]: row[1] for row in workbook["Progress"].iter_rows(values_only=True)}
    assert metrics["variance_basis"] == "baseline"


def test_ms_project_carries_actuals_and_baseline(api_client, tracked):
    from xml.etree import ElementTree as ET

    from app.exports.msproject import MSPDI_NS

    project, first, second = tracked
    root = ET.fromstring(api_client.get(f"/api/projects/{project}/export/mspdi").content)
    tasks = {
        node.findtext(f"{{{MSPDI_NS}}}Name"): node
        for node in root.findall(f".//{{{MSPDI_NS}}}Task")
    }
    done = tasks[first["label"]]
    assert done.findtext(f"{{{MSPDI_NS}}}PercentComplete") == "100"
    assert done.findtext(f"{{{MSPDI_NS}}}ActualFinish").startswith("2026-09-07")
    baseline = done.find(f"{{{MSPDI_NS}}}Baseline")
    assert baseline is not None
    assert baseline.findtext(f"{{{MSPDI_NS}}}Finish").startswith(first["finish_date"])
    assert tasks[second["label"]].findtext(f"{{{MSPDI_NS}}}PercentComplete") == "50"


def test_p6_carries_status_and_actuals(api_client, tracked):
    project, first, second = tracked
    lines = api_client.get(f"/api/projects/{project}/export/xer").text.splitlines()
    start = lines.index("%T\tTASK")
    fields = lines[start + 1].split("\t")[1:]
    rows = []
    for line in lines[start + 2 :]:
        if not line.startswith("%R"):
            break
        rows.append(dict(zip(fields, line.split("\t")[1:], strict=True)))
    by_name = {row["task_name"]: row for row in rows}
    assert by_name[first["label"]]["status_code"] == "TK_Complete"
    assert by_name[first["label"]]["act_end_date"].startswith("2026-09-07")
    assert by_name[second["label"]]["status_code"] == "TK_Active"
    assert float(by_name[second["label"]]["phys_complete_pct"]) == 50.0
    assert by_name[second["label"]]["remain_drtn_hr_cnt"] != by_name[second["label"]][
        "target_drtn_hr_cnt"
    ]


def test_json_carries_progress_for_4d_replay(api_client, tracked):
    import json

    project, first, _ = tracked
    package = json.loads(api_client.get(f"/api/projects/{project}/export/json").content)
    task = next(t for t in package["tasks"] if t["id"] == first["id"])
    assert task["progress"]["status"] == "complete"
    assert task["progress"]["actual_finish"] == "2026-09-07"
    assert task["source_global_ids"]
    assert package["progress"]["summary"]["variance_basis"] == "baseline"
    assert package["progress"]["baseline"]["name"] == "Contract"
